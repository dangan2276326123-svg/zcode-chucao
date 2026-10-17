# -*- coding: utf-8 -*-
"""Tests for tools/make_thesis_docx.py.

These guard the two ways this generator could do silent damage:
  * a symbol that is mangled but not flagged -- the file would look finished
    while showing wrong math (test_symbol_* below, found by an actual run);
  * a marker that is dropped instead of carried into the document -- the
    "resolve before submission" list would shrink on its own (the tallies must
    equal what the sources really contain).
"""
import io
import os
import re
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))), 'tools'))
import make_thesis_docx as M  # noqa: E402


# ------------------------------------------------------------------ symbols
def test_greek_and_superscript_map_clean():
    s, clean = M.symbol_to_text(r'\delta_{fr}')
    assert s.startswith('\u03b4') and clean, s


def test_hat_and_dot_become_combining_marks():
    s, clean = M.symbol_to_text(r'\hat{\tau}')
    assert clean and '\u0302' in s and '\\' not in s, s
    s2, clean2 = M.symbol_to_text(r'\dot{e}')
    assert clean2 and '\u0307' in s2, s2


def test_nested_structure_is_flagged_not_garbled():
    """The regression this test exists for: `e^{-\\tau s}` used to come out as
    'e\u207b\u03c4 s}' -- braces left behind -- while still reporting itself as
    successfully converted, so nothing highlighted it."""
    s, clean = M.symbol_to_text(r'e^{-\tau s}')
    assert not clean, 'a leftover brace must be reported as unmapped: %r' % s


def test_unmapped_command_keeps_its_backslash_and_is_flagged():
    s, clean = M.symbol_to_text(r'\Bbbk')
    assert not clean and '\\' in s


# --------------------------------------------------------------------- emit
class _Font(object):
    def __init__(self):
        self.highlight_color = None
        self.size = None
        self.name = None
        self.superscript = False


class _Run(object):
    """Minimal stand-in for a python-docx run: emit() touches .bold, .italic
    and .font.highlight_color, so the fake has to offer the same surface."""

    def __init__(self, text):
        self.text = text
        self.bold = False
        self.italic = False
        self.font = _Font()


class FakePara(object):
    """Stands in for a python-docx paragraph so emit() can be tested without
    writing a package."""

    def __init__(self):
        self.runs = []

    def add_run(self, text):
        r = _Run(text)
        self.runs.append(r)
        return r

    def plain(self):
        return ''.join(r.text for r in self.runs)


def test_emit_keeps_every_payload_visible():
    c = M.Counter()
    p = FakePara()
    M.emit(r'见[[SYM:\delta]]与[[EQ:x=\frac{a}{b}]]，另<待补充：台架实测>。'
           r'文献[[REF:Zhang 2000.]]', p, c)
    body = p.plain()
    assert c['sym_ok'] == 1 and c['eq'] == 1 and c['todo'] == 1
    assert c['ref_marker'] == 1
    # nothing may vanish: the formula text, the todo text and the citation text
    # all have to survive into the rendered runs
    assert r'\frac{a}{b}' in body
    assert '台架实测' in body and 'Zhang 2000' in body
    assert '见' in body and '。' in body


def test_resolved_citation_becomes_a_superscript_number():
    """[[REF:12]] 是**已定号**的引文：排成上标 [12]，不再进灰底占位那一桶。

    两栏计数必须分开——"还剩几处没定号"不能藏在"41 处标记"这种混合读数里。
    """
    c = M.Counter()
    p = FakePara()
    M.emit('如文献[[REF:12]]所示。', p, c)
    assert c['ref_marker'] == 1 and c['ref_num'] == 1 and c['ref_pending'] == 0
    assert '[12]' in p.plain() and '待补引文' not in p.plain()
    run = [r for r in p.runs if r.text == '[12]'][0]
    assert run.font.superscript is True, '顺序编码制的引用标记是上标'
    # 没定号的仍然走灰底占位，一条不许静默消失
    c2 = M.Counter()
    p2 = FakePara()
    M.emit('如[[REF:KEYWORD_PLACEHOLDER: Someone 2020.]]所示。', p2, c2)
    assert c2['ref_pending'] == 1 and c2['ref_num'] == 0
    assert '待补引文' in p2.plain()


def test_other_pending_markers_are_highlighted_too():
    """清零清单不只 `<待补充>`：正文里还有 `<待核实>`、`<待登记>`、`<待实测>`。

    原先只认"待补充"，那几处就以普通文字落纸——不高亮、也不进计数。
    清单少报比格式不好看更糟，因为这份文档里的 `<待…>` 计数就是送审前的清零对象。
    """
    c = M.Counter()
    p = FakePara()
    M.emit(u'见<待核实：HRNet 文献>与<待登记：轮盘外径>与<待实测：X>。', p, c)
    assert c['todo'] == 3, c['todo']
    marked = [r for r in p.runs if r.text.startswith(u'<待')]
    assert len(marked) == 3
    assert all(r.font.highlight_color is not None for r in marked), \
        '一处待办标记都不许以普通文字混过版面'


def test_bold_mark_produces_a_bold_run_and_no_asterisks():
    c = M.Counter()
    p = FakePara()
    M.emit('实际生效的是**比例控制加时延超前**（未整定）', p, c)
    assert '*' not in p.plain()
    assert any(r.bold and '比例控制' in r.text for r in p.runs)


def test_check_pass_counts_the_same_as_a_write():
    """add_table() used to return early in check mode, so table-cell markers
    were invisible to --check and its tallies quietly disagreed with a real
    run (seen as 27 vs 31)."""
    rows = [['指标', '值'], ['mIoU', '<待补充：独立测试集>']]
    c = M.Counter()
    M.add_table(None, rows, c, check_only=True)
    assert c['todo'] == 1, 'a <待补充> inside a table cell must still be tallied'


# ------------------------------------------------------------------ parsing
def test_parse_blocks_drops_separator_rows_and_keeps_cells():
    lines = ['表 2-2 参数', '', '| 参数 | 数值 |', '| --- | --- |',
             '| 车宽 | 2.15 m |', '']
    blocks = list(M.parse_blocks(lines))
    tables = [b for b in blocks if b[0] == 'table']
    assert len(tables) == 1
    assert tables[0][1] == [['参数', '数值'], ['车宽', '2.15 m']]


def test_parse_blocks_reads_heading_levels():
    blocks = list(M.parse_blocks(['# 章', '## 节', '### 小节', '正文一句']))
    assert [(b[0], b[1]) for b in blocks] == [('h', 1), ('h', 2), ('h', 3),
                                              ('p', '正文一句')]


# ------------------------------------------------- cross-check against source
def test_tallies_match_the_real_chapter_sources():
    """End-to-end guard: what the generator counts must equal what a plain
    scan of the .md sources contains. If a construct ever stops being parsed,
    this goes red instead of the document silently losing content."""
    if not os.path.isdir(M.CHAPTERS):
        pytest.skip('thesis chapter sources not on this machine')
    raw = {'todo': 0, 'ref': 0, 'eq': 0, 'fig': 0, 'sym': 0}
    for name in M.SOURCES:
        text = '\n'.join(M.read_source(name))
        raw['todo'] += len(re.findall(r'<待[^\n<>]{0,200}>', text, re.S))
        raw['ref'] += len(re.findall(r'\[\[REF:', text))
        raw['eq'] += len(re.findall(r'\[\[EQ:', text))
        raw['fig'] += len(re.findall(r'\[\[FIG:', text))
        raw['sym'] += len(re.findall(r'\[\[SYM:', text))
    c, _refs = M.build(check_only=True)
    assert c['todo'] == raw['todo'], (c['todo'], raw['todo'])
    assert c['ref_marker'] == raw['ref']
    assert c['eq'] == raw['eq']
    assert c['fig'] == raw['fig']
    assert c['sym_ok'] + c['sym_raw'] == raw['sym']


# --------------------------------------------------------------------- cover
# The cover is carried over as template XML rather than retyped, so these tests
# exist to catch the failure modes that are invisible on screen: a field left
# behind, the dead external picture link surviving, or a relationship id that no
# longer resolves.
needs_template = pytest.mark.skipif(
    not os.path.exists(M.TEMPLATE), reason='school template not on this machine')


def _blocks():
    blocks, tally = M.extract_front_matter(M.TEMPLATE, keep_box='academic')
    assert blocks, 'front matter must not come back empty'
    from lxml import etree
    return ''.join(etree.tostring(b).decode('utf-8') for b in blocks), tally


@needs_template
def test_carried_cover_has_no_field_machinery():
    xml, tally = _blocks()
    for tag in ('w:fldChar', 'w:instrText', 'MACROBUTTON', 'INCLUDEPICTURE',
                'w:sdt', 'w:altChunk'):
        assert tag not in xml, '%s survived the strip' % tag
    assert tally['fields_stripped'] > 0, 'the template cover does contain fields'


@needs_template
def test_dead_external_picture_link_is_dropped():
    """The template's cover INCLUDEPICTURE points at a 2005 news.cau.edu.cn jpg
    that now redirects to the university homepage; leaving it means a red X."""
    xml, tally = _blocks()
    assert 'r:link=' not in xml
    assert 'news.cau.edu.cn' not in xml
    assert tally['external_links_dropped'] >= 1


@needs_template
def test_local_wordmark_survives_and_every_rid_resolves():
    """Dropping the link must not drop the picture, and no carried block may
    reference a relationship the output package does not have."""
    import re
    import zipfile
    xml, _ = _blocks()
    assert 'v:imagedata' in xml, 'the school wordmark must still be there'
    rels = zipfile.ZipFile(M.TEMPLATE).read('word/_rels/document.xml.rels').decode('utf-8')
    have = set(re.findall(r'Id="([^"]+)"', rels))
    used = set(re.findall(r'r:(?:id|embed|href)="([^"]+)"', xml))
    assert not (used - have), 'dangling relationship ids: %s' % sorted(used - have)


@needs_template
def test_academic_choice_keeps_box1_and_drops_box2():
    """Text of the kept blocks: 附件1 stays (with the template's own
    "delete this box" instruction), 附件2 goes."""
    blocks, _ = M.extract_front_matter(M.TEMPLATE, keep_box='academic')
    from lxml import etree
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    text = '\n'.join(''.join(t.text or '' for t in b.iter(W + 't'))
                     for b in blocks)
    assert '附件1' in text
    assert '附件2' not in text
    assert '独 创 性 声 明' in text or '独创性声明' in text.replace(' ', '')


@needs_template
def test_seam_between_cover_and_spine_is_page_broken():
    """The template keeps its only explicit page break at the end of 附件2, so
    dropping that box would pull the 书脊 page up onto the cover."""
    xml, tally = _blocks()
    assert 'w:type="page"' in xml or tally['page_breaks_added'] >= 1


# ------------------------------------------------------------------ figures
def test_figure_map_splits_assets_from_placeholders(tmp_path):
    """清单里 File 有值 -> 嵌图；写 em-dash -> 留占位。

    这张表是论文插图的唯一事实源，所以两种状态必须分得开：
    「登记了但文件不在」是清单在说谎，不能和「还没做」混成同一个计数。
    """
    man = tmp_path / 'figures_manifest.md'
    bt = chr(96)                     # 反引号，清单里用它包文件名与标记
    rows = [
        (u'图3-2', bt + 'a.png' + bt, u'训练指标三段走势'),
        (u'图2-1', u'—', u'系统总体架构图'),
    ]
    man.write_text(
        u'# m' + chr(10) + chr(10)
        + u'| ID | File | 正文里的占位标记 | Related Chapter | Source | Status | Notes |'
        + chr(10)
        + u'| --- | --- | --- | --- | --- | --- | --- |' + chr(10)
        + chr(10).join(u'| %s | %s | %s[[FIG:%s]]%s | ch3 | x | ok | y |'
                       % (r[0], r[1], bt, r[2], bt) for r in rows) + chr(10),
        encoding='utf-8')
    m = M.load_figure_map(str(man))
    assert m[u'训练指标三段走势'] == 'a.png'
    assert m[u'系统总体架构图'] is None


def test_placeholder_is_kept_when_no_asset(monkeypatch):
    """没有资产时标记必须留在正文里，不能静默消失——标记一丢，
    「这张图还没做」这件事就从正文和清单两边同时蒸发了。"""
    monkeypatch.setattr(M, 'FIGURE_MAP', {u'某图': None})
    c = M.Counter()
    p = FakePara()
    M.emit(u'见图[[FIG:某图]]。', p, c)
    assert c['fig_placeholder'] == 1
    assert u'待补图' in p.plain()


def test_registered_but_missing_file_is_counted_separately(monkeypatch, tmp_path):
    """清单登记了文件名、盘上却没有 -> 单独计数，不与「还没做」合并。"""
    monkeypatch.setattr(M, 'FIGURE_MAP', {u'某图': 'no_such_file_anywhere.png'})
    monkeypatch.setattr(M, 'FIGDIR', str(tmp_path))   # 空目录：登记了但没有文件
    c = M.Counter()
    M.emit(u'[[FIG:某图]]', FakePara(), c)
    assert c['fig_missing_file'] == 1
    assert c['fig_placeholder'] == 0


def test_embedded_figure_paragraph_drops_exact_line_spacing(tmp_path, monkeypatch):
    """嵌图那一段必须离开"固定值"行距。

    这是 v15 导出 PDF 才看见的缺陷：正文样式是固定值行距，Word 会把嵌入图片裁到
    一行高——两张 9×16 cm 的板卡照片在页面上各只推进 36 pt，图块还互相重叠。
    计数、`--verify` 的域读数、`Fields.Count` 全都看不出这件事，
    所以这条测试按 XML 判：段间距的 lineRule 不许还是 exact。
    """
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from PIL import Image

    png = tmp_path / 'f.png'
    Image.new('RGB', (600, 200), 'white').save(str(png))
    monkeypatch.setattr(M, 'FIGURE_MAP', {u'某图': 'f.png'})
    monkeypatch.setattr(M, 'FIGDIR', str(tmp_path))

    doc = Document()
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(12)          # 模板正文样式就是这样
    assert p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY, \
        '这条测试的前提要先成立：先做出固定值行距，才谈得上被改掉'
    # 校模板的缩进是**按字符**给的（样式里 w:firstLineChars="200" 有 4 处），
    # 只清 w:firstLine 顶不掉它——第一版就是这么漏的，图仍被推右半个字。
    p._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'), '200')

    c = M.Counter()
    M.emit(u'[[FIG:某图]]', p, c)
    assert c['fig_embedded'] == 1

    spacing = p._p.get_or_add_pPr().find(qn('w:spacing'))
    assert spacing is not None and spacing.get(qn('w:lineRule')) == 'auto', \
        '图段仍是固定值行距，Word 会把图片压成一行高'
    assert p.paragraph_format.line_spacing == 1
    ind = p._p.get_or_add_pPr().find(qn('w:ind'))
    assert p.paragraph_format.first_line_indent.cm == 0, \
        '图段带首行缩进会把居中的图推偏半个字'
    assert ind.get(qn('w:firstLineChars')) == '0', \
        '字符单位的缩进没清：w:firstLine="0" 顶不过 w:firstLineChars'


# ---------------------------------------------------------------- page layout
def test_level_one_heading_starts_a_new_page():
    """校模板的"标题 1"没有段前分页（styles.xml 里 pageBreakBefore 出现 0 次），
    所以章标题会接在上一章末尾——学校格式要求每章另起一页。"""
    from docx import Document
    doc = Document()
    r = M.Renderer(doc, M.Counter())
    r.heading(1, u'第 1 章 绪论')
    r.heading(2, u'1.1 研究背景')
    assert doc.paragraphs[-2].paragraph_format.page_break_before is True
    assert not doc.paragraphs[-1].paragraph_format.page_break_before, \
        '二级标题不该另起一页'


needs_prev = pytest.mark.skipif(
    not os.path.exists(M.PREV_DOCX), reason='previous docx not on this machine')


@needs_template
@needs_prev
def test_body_section_carries_running_header_and_page_numbers(tmp_path, monkeypatch):
    """整份生成一次，验正文那一节的页眉与两节的页码。

    这条只能整份建：症状来自"沿用模板最后一节的页面设置"，而模板最后一节是作者简介
    ——它的 default 页眉写死"作者简介"和"博/硕士"待选文本，真正的正文页眉
    （右格 STYLEREF 标题 1）挂在 even 上，而本包没开奇偶页，于是每一页右上角都是
    "作者简介"。域计数、Fields.Count、嵌图计数全都看不出这件事。
    """
    import re

    from docx import Document
    from docx.oxml.ns import qn

    monkeypatch.setattr(M, 'OUTDIR', str(tmp_path))     # 绝不往 thesis_cau 里写
    monkeypatch.setattr(M, 'OUT_DOCX', str(tmp_path / 'built.docx'))
    c, _ = M.build(check_only=False)
    assert c['hdr_repointed'] == 1, '没找到可改向的跑动页眉，改动就没生效'

    d = Document(str(tmp_path / 'built.docx'))
    front, body = d.sections[-2], d.sections[-1]

    def default_rid(section, tag):
        rids = [r.get(qn('r:id')) for r in section._sectPr.findall(qn(tag))
                if (r.get(qn('w:type')) or 'default') == 'default']
        assert len(rids) == 1, '%s 的 default %s 应有且只有一条：%s' % (
            'front' if section is front else 'body', tag, rids)
        return rids[0]

    blob = d.part.related_parts[default_rid(body, 'w:headerReference')].blob
    assert b'STYLEREF' in blob, '正文页眉没指向跑动标题部件'
    assert u'作者简介'.encode('utf-8') not in blob, '页眉仍是作者简介那一份'
    assert u'双击页眉'.encode('utf-8') not in blob, \
        '挑中了模板的占位页眉——那会把"双击页眉修改此处章节名称"印在每一页上'

    # 页码：分节之后前置走罗马数字、正文从 1 起阿拉伯数字，两节都挂带 PAGE 域的页脚。
    assert c['page_front'] == 'upperRoman@1' and c['page_body'] == 'decimal@1'
    for name, section in (('front', front), ('body', body)):
        pg = section._sectPr.find(qn('w:pgNumType'))
        assert pg is not None and pg.get(qn('w:start')) == '1', name
        foot = d.part.related_parts[default_rid(section, 'w:footerReference')].blob
        assert re.search(rb'\bPAGE\b', foot), '%s 页脚没挂 PAGE' % name
    assert d.part.related_parts[default_rid(front, 'w:headerReference')].blob \
        == d.part.related_parts[default_rid(d.sections[0], 'w:headerReference')].blob, \
        '前置页眉应沿用封面那份空的，否则摘要头上会顶着跑动章名'
    # 正文那一节是 add_section 造出来的空节：页眉页脚都得显式挂上，缺一即回归。
    assert default_rid(body, 'w:footerReference') != default_rid(body, 'w:headerReference')


# ------------------------------------------------------------------ 参考文献订正
def _fake_refs():
    out = ['[%d] filler entry %d.' % (i, i) for i in range(1, 30)]
    out[3] = ("[4] Ruckelshausen A, et al. BoniRob—an autonomous field robot "
              "platform[C]//Precision Agriculture '09. Wageningen, 2009: 741-747.")
    out[15] = ("[16] Chen L C, et al. Encoder-decoder with atrous separable "
               "convolution[C]//ECCV. Munich: Springer, 2018: 801-818.")
    return out


def test_reference_fixes_apply_and_are_counted():
    refs, n = M.apply_reference_fixes(_fake_refs())
    assert n == len(M.REFERENCE_FIXES) == 2
    assert '841-847' in refs[3] and '741-747' not in refs[3]
    assert '833-851' in refs[15] and '801-818' not in refs[15]


def test_reference_fix_bites_when_the_list_changes():
    """订正**不许静默失效**：表换版或有人手改过，原串找不到就必须抛。

    否则错页码会悄悄回到正文里，而输出照旧打印"26 条搬运完成"。
    """
    refs = _fake_refs()
    refs[3] = refs[3].replace('741-747', '999-999')
    with pytest.raises(ValueError) as e:
        M.apply_reference_fixes(refs)
    assert '[4]' in str(e.value)

    short = ['[1] only one entry.']
    with pytest.raises(ValueError):
        M.apply_reference_fixes(short)
